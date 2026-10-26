# -*- coding: utf-8 -*-
"""
文書監査の報告DOCXを、正本から生成する。

正本（これを直す。DOCXは直さない）:
  audit/confirmations.json      … 実行者への確認事項。IDで継続追跡する
  audit/findings/<日付>.md      … 回次ごとの監査記録

生成物:
  audit/dist/markdown_audit_<YYYYMMDDHHMMSS>.docx

同じ監査を再実行しても、DOCXは**日時付きの別ファイル**として出力され、既存ファイルを
上書きしない。内容は毎回正本から作り直すため、同一事項が重複して積み上がることはない。

使い方:

    python tools/gen_audit_docx.py                 # 最新回の findings で生成
    python tools/gen_audit_docx.py --date 2026-08-04
    python tools/gen_audit_docx.py --out <パス>    # 出力先を明示する
    python tools/gen_audit_docx.py --list          # 登録済み確認事項の一覧だけ表示

依存: python-docx（標準ライブラリ以外はこれだけ）
"""
import argparse
import datetime
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
AUDIT = os.path.join(ROOT, "audit")
CONFIRMATIONS = os.path.join(AUDIT, "confirmations.json")
FINDINGS_DIR = os.path.join(AUDIT, "findings")
DIST = os.path.join(AUDIT, "dist")

PROJECT_NAME = "make_automate_flow — Power Automate開発基盤"
DOC_TITLE = "Markdown・Skill整合性監査報告書"

# 確認事項の表に出す列（依頼仕様の12項目）
CONF_COLUMNS = [
    ("id", "ID"),
    ("category", "分類"),
    ("target", "対象"),
    ("question", "確認したいこと"),
    ("reason", "確認が必要な理由"),
    ("steps", "実行手順"),
    ("expected", "期待される結果"),
    ("returnToClaude", "Claudeへ返してほしい情報"),
    ("priority", "優先度"),
    ("status", "状態"),
    ("foundOn", "発見日"),
    ("relatedFinding", "関連する不整合"),
]


def load_confirmations():
    with open(CONFIRMATIONS, encoding="utf-8") as fh:
        return json.load(fh)


def latest_findings_date():
    if not os.path.isdir(FINDINGS_DIR):
        return None
    dates = sorted(f[:-3] for f in os.listdir(FINDINGS_DIR)
                   if re.match(r"^\d{4}-\d{2}-\d{2}\.md$", f))
    return dates[-1] if dates else None


def parse_findings(path):
    """findings/<日付>.md を見出し単位の辞書に分解する。

    DOCXの章立ては本関数の戻り値から組み立てる。Markdown側の見出しを変えたら
    SECTION_MAP を合わせる（正本はあくまで Markdown 側）。
    """
    with open(path, encoding="utf-8") as fh:
        text = fh.read()
    sections, current, buf = {}, None, []
    for line in text.splitlines():
        m = re.match(r"^## +(.*)$", line)
        if m:
            if current:
                sections[current] = "\n".join(buf).strip()
            current, buf = m.group(1).strip(), []
        else:
            buf.append(line)
    if current:
        sections[current] = "\n".join(buf).strip()
    return sections


def find_section(sections, *keywords):
    """見出しにキーワードを含む節を返す（見出しの表記ゆれに耐えるため）"""
    for head, body in sections.items():
        if any(k in head for k in keywords):
            return body
    return ""


# =========================================================================
# Markdown の簡易レンダリング（見出し・表・箇条書き・段落）
# =========================================================================
def strip_inline(s):
    """DOCXは等幅・太字装飾を持ち込まないので、Markdownの装飾記号を落とす"""
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)   # リンクはテキストだけ残す
    s = s.replace("**", "").replace("`", "")
    return s.strip()


def add_md_block(doc, md, level=2):
    """Markdown断片を段落・箇条書き・表として本文へ流し込む"""
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue

        # 表（| a | b |）
        if line.lstrip().startswith("|") and i + 1 < len(lines) and \
                re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                if not re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i]):
                    cells = [strip_inline(c) for c in lines[i].strip().strip("|").split("|")]
                    rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows[0], rows[1:])
            continue

        # 見出し（### 以下）
        m = re.match(r"^(#{3,6}) +(.*)$", line)
        if m:
            doc.add_heading(strip_inline(m.group(2)), level=min(level + 1, 4))
            i += 1
            continue

        # 箇条書き / 番号付き
        m = re.match(r"^\s*[-*] +(.*)$", line)
        if m:
            doc.add_paragraph(strip_inline(m.group(1)), style="List Bullet")
            i += 1
            continue
        m = re.match(r"^\s*\d+\. +(.*)$", line)
        if m:
            doc.add_paragraph(strip_inline(m.group(1)), style="List Number")
            i += 1
            continue

        # コードブロックは等幅で1段落にまとめる
        if line.lstrip().startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].lstrip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            if code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            continue

        # 引用
        if line.lstrip().startswith(">"):
            doc.add_paragraph(strip_inline(line.lstrip("> ")), style="Intense Quote")
            i += 1
            continue

        doc.add_paragraph(strip_inline(line))
        i += 1


def add_table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for j, h in enumerate(header):
        cell = t.rows[0].cells[j]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for j in range(len(header)):
            cells[j].text = row[j] if j < len(row) else ""
    return t


# =========================================================================
def build(doc, conf, findings_date, sections):
    items = conf["items"]

    # ---- 1. 表紙 ----
    doc.add_paragraph()
    h = doc.add_heading(DOC_TITLE, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(PROJECT_NAME)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    add_table(doc, ["項目", "内容"], [
        ["プロジェクト名", PROJECT_NAME],
        ["文書名", DOC_TITLE],
        ["監査実施日", findings_date],
        ["対象時点", "Git管理外のため、作業ディレクトリの現物（D:\\make_automate_flow）"],
        ["正本データ", "audit/confirmations.json ／ audit/findings/%s.md" % findings_date],
        ["生成方法", "python tools/gen_audit_docx.py（本書は自動生成。直接編集しない）"],
        ["生成日時", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ])
    doc.add_page_break()

    # ---- 2〜6. findings の各節 ----
    doc.add_heading("1. 今回の監査概要", level=1)
    add_md_block(doc, find_section(sections, "監査概要"))

    doc.add_heading("2. 今回確認した範囲", level=1)
    add_md_block(doc, find_section(sections, "確認した範囲"))

    doc.add_heading("3. 発見した主な不整合", level=1)
    add_md_block(doc, find_section(sections, "発見した不整合", "不整合と対応"))

    doc.add_heading("4. 実際に修正した内容", level=1)
    add_md_block(doc, find_section(sections, "実際に修正した内容"))

    doc.add_heading("5. Skill、Agent及びRuleの修正内容", level=1)
    add_md_block(doc, find_section(sections, "Skill・Agent の見直し", "Skill"))

    # ---- 7. 確認事項（正本 JSON から生成） ----
    doc.add_page_break()
    doc.add_heading("6. 実行者への確認事項", level=1)
    open_items = [i for i in items if i["status"] in ("未確認", "確認待ち")]
    doc.add_paragraph(
        "登録 %d 件（未対応 %d 件）。正本は audit/confirmations.json。"
        "確認済みの項目も削除せず、状態と確認日を記録して継続追跡する。"
        % (len(items), len(open_items)))

    add_table(doc, ["ID", "分類", "対象", "優先度", "状態"],
              [[i["id"], i["category"], i["target"], i["priority"], i["status"]] for i in items])

    for it in items:
        doc.add_heading("%s（%s・優先度%s）" % (it["id"], it["category"], it["priority"]), level=2)
        rows = []
        for key, label in CONF_COLUMNS:
            if key in ("id", "category", "priority"):
                continue
            val = it.get(key)
            if isinstance(val, list):
                val = "\n".join("%d. %s" % (n, s) for n, s in enumerate(val, 1))
            rows.append([label, "" if val is None else str(val)])
        if it.get("notes"):
            rows.append(["補足", it["notes"]])
        add_table(doc, ["項目", "内容"], rows)
        doc.add_paragraph()

    # ---- 8〜9 ----
    doc.add_heading("7. 未解決事項", level=1)
    add_md_block(doc, find_section(sections, "未解決事項"))

    doc.add_heading("8. 次回監査時の確認ポイント", level=1)
    add_md_block(doc, find_section(sections, "次回監査"))

    # ---- 10. 変更ファイル一覧 ----
    doc.add_heading("9. 変更ファイル一覧", level=1)
    changed = find_section(sections, "実際に修正した内容")
    m = re.search(r"###? *変更ファイル一覧\s*(.+)$", changed, re.S)
    add_md_block(doc, m.group(1).strip() if m else changed)


def main():
    ap = argparse.ArgumentParser(description="監査の正本から報告DOCXを生成する")
    ap.add_argument("--date", help="対象の監査回（audit/findings/<日付>.md）。既定は最新")
    ap.add_argument("--out", help="出力先パス。既定は audit/dist/markdown_audit_<日時>.docx")
    ap.add_argument("--list", action="store_true", help="確認事項の一覧を表示して終了")
    args = ap.parse_args()

    if not os.path.isfile(CONFIRMATIONS):
        print("正本が見つからない: %s" % CONFIRMATIONS)
        return 1
    conf = load_confirmations()

    if args.list:
        print("確認事項 %d 件（正本: audit/confirmations.json）" % len(conf["items"]))
        for i in conf["items"]:
            print("  %-9s %-6s %-4s %s" % (i["id"], i["status"], i["priority"], i["question"][:56]))
        return 0

    date = args.date or latest_findings_date()
    if not date:
        print("監査記録が無い: %s/<日付>.md を先に作成する" % FINDINGS_DIR)
        return 1
    fpath = os.path.join(FINDINGS_DIR, "%s.md" % date)
    if not os.path.isfile(fpath):
        print("監査記録が見つからない: %s" % fpath)
        return 1

    try:
        global Document, Pt, WD_ALIGN_PARAGRAPH
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        print("python-docx が必要: pip install python-docx")
        return 1

    sections = parse_findings(fpath)
    doc = Document()
    build(doc, conf, date, sections)

    os.makedirs(DIST, exist_ok=True)
    out = args.out or os.path.join(
        DIST, "markdown_audit_%s.docx" % datetime.datetime.now().strftime("%Y%m%d%H%M%S"))
    if os.path.exists(out):
        print("既存ファイルを上書きしない: %s" % out)
        return 1
    doc.save(out)

    print("生成: %s" % os.path.relpath(out, ROOT).replace(os.sep, "/"))
    print("  監査回      : %s" % date)
    print("  確認事項    : %d 件（未対応 %d 件）"
          % (len(conf["items"]),
             len([i for i in conf["items"] if i["status"] in ("未確認", "確認待ち")])))
    return 0


if __name__ == "__main__":
    sys.exit(main())
